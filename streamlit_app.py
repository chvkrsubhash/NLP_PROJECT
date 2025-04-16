import streamlit as st
import re
import io
import base64
import os
import random
from datetime import datetime
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
import pyrebase
from textblob import TextBlob, Word
try:
    import google.generativeai as genai
except ImportError:
    genai = None
import json

# Must be first
st.set_page_config(page_title="Technical Interview Chatbot", layout="wide")

# Dependencies
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
    st.error("PyPDF2 missing. Install with: pip install PyPDF2")

try:
    import docx
except ImportError:
    docx = None
    st.error("python-docx missing. Install with: pip install python-docx")

from fpdf import FPDF

# Firebase configuration
firebase_config = {
    "apiKey": "AIzaSyDvFFLr-Fjhma2yae7rx3r7Ei0J6bXJmmI",
    "authDomain": "client-2bbfc.firebaseapp.com",
    "databaseURL": "https://client-2bbfc-default-rtdb.firebaseio.com",
    "projectId": "client-2bbfc",
    "storageBucket": "client-2bbfc.firebasestorage.app",
    "messagingSenderId": "971318119261",
    "appId": "1:971318119261:web:0cf9b5f290f1589326f6b4",
    "measurementId": "G-5YHQGXBXJG"
}

# Initialize Firebase
try:
    firebase = pyrebase.initialize_app(firebase_config)
    auth = firebase.auth()
except Exception as e:
    st.error(f"Firebase initialization failed: {e}. Check secrets.")
    auth = None

# Email configuration
EMAIL_SENDER = "projecttestingsubhash@gmail.com"
EMAIL_PASSWORD = "zgwynxksfnwzusyk"
SMTP_SERVER = "smtp.gmail.com"
SMTP_PORT =  587

# Skill definitions
COMMON_SKILLS = {
    'programming': ['python', 'java', 'javascript', 'html', 'css', 'c++', 'c#', 'ruby', 'php', 'sql'],
    'frameworks': ['react', 'angular', 'vue', 'django', 'flask', 'spring', 'node.js', 'express', '.net'],
    'databases': ['sql', 'mysql', 'postgresql', 'mongodb', 'oracle', 'sqlite', 'redis'],
    'cloud': ['aws', 'azure', 'gcp', 'docker', 'kubernetes'],
    'tools': ['git', 'github', 'jira', 'jenkins', 'agile', 'scrum'],
}

# Question bank
TECHNICAL_QUESTIONS = {
    'python': [
        {"question": "Explain Python decorators with an example.", 
         "expected_keywords": ["function", "wrapper", "decorator", "@", "arguments", "return"]},
        {"question": "How do you handle exceptions in Python?", 
         "expected_keywords": ["try", "except", "finally", "raise", "error", "handling"]},
        {"question": "What are list comprehensions in Python?", 
         "expected_keywords": ["list", "comprehension", "concise", "loop", "condition", "performance"]},
        {"question": "Explain Python generators and their benefits.", 
         "expected_keywords": ["generator", "yield", "iterator", "memory", "lazy", "evaluation"]},
        {"question": "How does Python manage memory?", 
         "expected_keywords": ["garbage collection", "reference", "counting", "memory", "del"]},
    ],
    'java': [
        {"question": "What is inheritance in Java?", 
         "expected_keywords": ["extends", "class", "parent", "child", "super", "override"]},
        {"question": "How does Java handle memory management?", 
         "expected_keywords": ["garbage collection", "heap", "stack", "reference", "finalize"]},
        {"question": "Explain Java streams API.", 
         "expected_keywords": ["stream", "functional", "map", "filter", "collect", "pipeline"]},
        {"question": "What is synchronization in Java threads?", 
         "expected_keywords": ["synchronized", "thread", "lock", "monitor", "concurrency"]},
        {"question": "Describe Java's Optional class.", 
         "expected_keywords": ["optional", "null", "avoid", "check", "orElse", "present"]},
    ],
    'javascript': [
        {"question": "What are closures in JavaScript?", 
         "expected_keywords": ["function", "scope", "variable", "closure", "lexical", "access"]},
        {"question": "Explain event delegation in JavaScript.", 
         "expected_keywords": ["event", "delegation", "bubble", "target", "listener", "parent"]},
        {"question": "What are promises and async/await?", 
         "expected_keywords": ["promise", "async", "await", "resolve", "reject", "asynchronous"]},
        {"question": "How does the JavaScript event loop work?", 
         "expected_keywords": ["event loop", "call stack", "queue", "async", "callback"]},
        {"question": "Difference between let, const, and var?", 
         "expected_keywords": ["scope", "let", "const", "var", "block", "hoisting"]},
    ],
    'sql': [
        {"question": "What's the difference between INNER and LEFT JOIN?", 
         "expected_keywords": ["inner", "left", "join", "matching", "all", "records"]},
        {"question": "How do you optimize a slow SQL query?", 
         "expected_keywords": ["index", "execution plan", "query", "optimize", "performance"]},
        {"question": "What are SQL triggers?", 
         "expected_keywords": ["trigger", "event", "table", "insert", "update", "delete"]},
        {"question": "Explain ACID properties.", 
         "expected_keywords": ["atomicity", "consistency", "isolation", "durability", "transaction"]},
        {"question": "What is a CTE in SQL?", 
         "expected_keywords": ["common table expression", "with", "query", "temporary", "recursive"]},
    ],
    'aws': [
        {"question": "What's the difference between EC2 and Lambda?", 
         "expected_keywords": ["instance", "serverless", "EC2", "Lambda", "scaling", "compute"]},
        {"question": "How do you secure an AWS environment?", 
         "expected_keywords": ["IAM", "security group", "encryption", "access", "policy"]},
        {"question": "What is AWS S3 used for?", 
         "expected_keywords": ["S3", "storage", "bucket", "object", "access", "policy"]},
        {"question": "Explain VPC components.", 
         "expected_keywords": ["VPC", "subnet", "route table", "gateway", "security group"]},
        {"question": "What is AWS CloudFormation?", 
         "expected_keywords": ["CloudFormation", "template", "infrastructure", "stack", "provision"]},
    ],
}

GENERIC_QUESTIONS = [
    {"question": "Describe a technical challenge you solved.", 
     "expected_keywords": ["challenge", "project", "solution", "overcome", "team", "result"]},
    {"question": "How do you stay updated with technology?", 
     "expected_keywords": ["learning", "research", "practice", "community", "courses"]},
    {"question": "What's your approach to debugging?", 
     "expected_keywords": ["debugging", "logs", "breakpoint", "systematic", "testing", "root cause"]},
    {"question": "How do you prioritize project tasks?", 
     "expected_keywords": ["prioritize", "deadline", "impact", "stakeholder", "planning"]},
    {"question": "Explain a time you improved a process.", 
     "expected_keywords": ["process", "improvement", "efficiency", "solution", "impact"]},
]

# Messages
WELCOME_MESSAGES = [
    "Welcome to TechInterviewBot! Let's practice your technical skills.",
    "Hello! Ready for a technical interview? I'm here to help.",
    "Hi! Let's sharpen your interview skills with tailored questions.",
]

RESUME_PROMPTS = [
    "Upload your resume or paste its content to start.",
    "Share your resume to tailor the interview questions.",
    "I need your resume to generate relevant questions.",
]

SKILL_MESSAGES = [
    "Great! Here are the skills I found in your resume:",
    "Thanks! I've identified these skills from your resume:",
    "Based on your resume, here are your key skills:",
]

INTERVIEW_START_MESSAGES = [
    "Let's start with questions based on your skills!",
    "Ready? Here come some technical questions!",
    "The interview begins with tailored questions.",
]

QUESTION_TRANSITIONS = [
    "Next question:",
    "Here's another one:",
    "Moving on:",
]

EVALUATION_POSITIVE = [
    "Great answer! You hit the key points.",
    "Excellent! Your response was clear and accurate.",
    "Well done! That was a strong answer.",
]

EVALUATION_AVERAGE = [
    "Good try! You covered some points, but there's room to grow.",
    "Decent answer, but you could add more detail.",
    "Not bad! Try expanding on the concepts.",
]

EVALUATION_NEEDS_IMPROVEMENT = [
    "You missed some key concepts. Let's review those.",
    "Needs more depth. Want some pointers?",
    "Try including more technical details.",
]

# Text sanitization for PDF
def sanitize_text(text):
    if not text:
        return ""
    replacements = {
        '\u2019': "'", '\u2018': "'", '\u201c': '"', '\u201d': '"', '\u2013': '-', '\u2014': '--',
    }
    for unicode_char, ascii_char in replacements.items():
        text = text.replace(unicode_char, ascii_char)
    return text.encode('ascii', 'ignore').decode('ascii')

# Email sending
def send_email(recipient_email, interview_record, pdf_path=None):
    if not recipient_email:
        st.error("No email address available. Please log in again.")
        return False
    try:
        msg = MIMEMultipart()
        msg['From'] = EMAIL_SENDER
        msg['To'] = recipient_email
        msg['Subject'] = f"Technical Interview Results - {interview_record['candidate_name']}"
        
        avg_score = interview_record.get('avg_score', 0)
        rating = interview_record.get('rating', 'N/A')
        skills = ", ".join([f"{category}: {', '.join(skills)}" for category, skills in interview_record.get('skills', {}).items()])
        
        body = f"""
        Dear {interview_record['candidate_name']},
        
        Your technical interview is complete!
        
        Overall Score: {avg_score:.1f}/100
        Rating: {rating}
        Skills: {skills if skills else 'None identified'}
        
        See the attached PDF for details (if available).
        
        Thank you,
        TechInterviewBot
        """
        msg.attach(MIMEText(body, 'plain'))
        
        if pdf_path and os.path.exists(pdf_path):
            with open(pdf_path, 'rb') as f:
                part = MIMEBase('application', 'octet-stream')
                part.set_payload(f.read())
            encoders.encode_base64(part)
            part.add_header(
                'Content-Disposition',
                f'attachment; filename={os.path.basename(pdf_path)}'
            )
            msg.attach(part)
        
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as server:
            server.starttls()
            server.login(EMAIL_SENDER, EMAIL_PASSWORD)
            server.send_message(msg)
        
        return True
    except Exception as e:
        st.error(f"Email sending failed: {e}")
        return False

# NLP functions using TextBlob
def preprocess_text(text):
    try:
        blob = TextBlob(text.lower())
        tokens = [Word(word).lemmatize() for word in blob.words if word.isalnum()]
        stop_words = {'a', 'an', 'and', 'are', 'as', 'at', 'be', 'by', 'for', 'from', 'has', 'he', 
                      'in', 'is', 'it', 'its', 'of', 'on', 'that', 'the', 'to', 'was', 'were', 'will', 'with'}
        tokens = [t for t in tokens if t not in stop_words]
        return " ".join(tokens)
    except Exception:
        return text.lower()

def extract_skills(text):
    if not text:
        return {}
    
    try:
        raw_text = text.lower()
        blob = TextBlob(raw_text)
        identified_skills = {}
        debug_matches = []
        
        context_patterns = r'(?:skills|experience|proficient in|worked with|knowledge of|using|expertise in)'
        
        for category, skill_list in COMMON_SKILLS.items():
            found_skills = set()
            for skill in skill_list:
                pattern = rf'{context_patterns}\s*[^.\n]*\b{re.escape(skill)}\b[^.\n]*'
                matches = re.finditer(pattern, raw_text)
                for match in matches:
                    context = match.group()
                    if skill in [Word(w).lemmatize() for w in TextBlob(context).words]:
                        found_skills.add(skill)
                        debug_matches.append(f"Matched '{skill}' in: '{context[:50]}...'")
            if found_skills:
                identified_skills[category] = list(found_skills)
        
        st.session_state.debug_skills = debug_matches
        st.session_state.raw_resume_text = raw_text
        return identified_skills
    except Exception:
        identified_skills = {}
        for category, skill_list in COMMON_SKILLS.items():
            found_skills = set()
            for skill in skill_list:
                pattern = rf'{context_patterns}\s*[^.\n]*\b{re.escape(skill)}\b[^.\n]*'
                if re.search(pattern, text.lower()):
                    found_skills.add(skill)
            if found_skills:
                identified_skills[category] = list(found_skills)
        return identified_skills

# Gemini-only evaluate_answer
def evaluate_answer(question, answer, expected_keywords):
    if not answer.strip():
        return {"score": 0, "feedback": "No answer provided.", "missing_concepts": expected_keywords}
    # if not genai:
    #     return {"score": 0, "feedback": "Gemini API not available.", "missing_concepts": expected_keywords}
    # GEMINI_API_KEY = "AIzaSyDlb0thGyUHOBuT5bmv9a8QCkg-UX5iMgY"
    # if not GEMINI_API_KEY:
    #     return {"score": 0, "feedback": "No Gemini API key.", "missing_concepts": expected_keywords}
    try:
        genai.configure(api_key="AIzaSyDlb0thGyUHOBuT5bmv9a8QCkg-UX5iMgY")
        model = genai.GenerativeModel('gemini-1.5-flash')
        prompt = f"""
        You are a technical interviewer. Evaluate:
        **Question**: {question}
        **Answer**: {answer}
        **Expected Concepts**: {', '.join(expected_keywords)}
        Return JSON:
        {{"score": <0-100>, "feedback": "<string>", "missing_concepts": [<string>, ...]}}
        """
        response = model.generate_content(prompt)
        if response and hasattr(response, 'text') and response.text:
            result = json.loads(response.text.strip().replace("```json\n", "").replace("\n```", ""))
            score = min(max(result.get("score", 0), 0), 100)
            return {"score": score, "feedback": result.get("feedback", "No feedback."), "missing_concepts": result.get("missing_concepts", [])}
        return {"score": 0, "feedback": "Invalid Gemini response.", "missing_concepts": expected_keywords}
    except Exception as e:
        return {"score": 0, "feedback": f"Gemini failed: {str(e)}", "missing_concepts": expected_keywords}

# File processing
def extract_text_from_pdf(pdf_file):
    if PyPDF2 is None:
        return ""
    try:
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        return "".join(page.extract_text() or "" for page in pdf_reader.pages)
    except Exception:
        return ""

def extract_text_from_docx(docx_file):
    if docx is None:
        return ""
    try:
        doc = docx.Document(docx_file)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception:
        return ""

# Question generation
def generate_technical_questions(skills, max_questions=5):
    all_possible_questions = []
    all_skills = [skill for category, skill_list in skills.items() for skill in skill_list]
    
    selected_skills = random.sample(all_skills, min(len(all_skills), 3)) if all_skills else []
    
    for skill in selected_skills:
        if skill in TECHNICAL_QUESTIONS:
            all_possible_questions.extend(random.sample(TECHNICAL_QUESTIONS[skill], min(2, len(TECHNICAL_QUESTIONS[skill]))))
    
    remaining_slots = max_questions - len(all_possible_questions)
    if remaining_slots > 0:
        all_possible_questions.extend(random.sample(GENERIC_QUESTIONS, min(remaining_slots, len(GENERIC_QUESTIONS))))
        remaining_slots = max_questions - len(all_possible_questions)
    
    if remaining_slots > 0:
        other_questions = [q for skill in TECHNICAL_QUESTIONS if skill not in selected_skills for q in TECHNICAL_QUESTIONS[skill]]
        all_possible_questions.extend(random.sample(other_questions, min(remaining_slots, len(other_questions))))
    
    random.shuffle(all_possible_questions)
    
    unique_questions = []
    question_texts = set()
    for q in all_possible_questions:
        if q["question"] not in question_texts:
            unique_questions.append(q)
            question_texts.add(q["question"])
    
    return unique_questions[:max_questions]

def get_feedback_message(score):
    if score >= 80:
        return random.choice(EVALUATION_POSITIVE)
    elif score >= 60:
        return random.choice(EVALUATION_AVERAGE)
    else:
        return random.choice(EVALUATION_NEEDS_IMPROVEMENT)

def format_skills_message(skills):
    return "\n".join(f"**{category.capitalize()}**: {', '.join(skill_list)}" for category, skill_list in skills.items())

# PDF export
def export_results_as_pdf(interview_record):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Technical Interview Results", ln=True, align="C")
    pdf.ln(5)
    
    candidate_name = sanitize_text(interview_record.get("candidate_name", "Candidate"))
    interview_date = interview_record.get("date", datetime.now().strftime("%Y-%m-%d %H:%M"))
    evaluations = interview_record.get("evaluations", {})
    total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
    avg_score = total_score / len(evaluations) if evaluations else 0
    rating = "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement"
    skills = interview_record.get("skills", {})
    questions = interview_record.get("questions", [])
    
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, f"Candidate: {candidate_name}", ln=True)
    pdf.cell(0, 10, f"Date: {interview_date}", ln=True)
    pdf.ln(5)
    
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Summary", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.cell(0, 10, f"Overall Score: {avg_score:.1f}/100", ln=True)
    pdf.cell(0, 10, f"Rating: {rating}", ln=True)
    pdf.ln(5)
    
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Skills", ln=True)
    pdf.set_font("Arial", "", 12)
    for category, skill_list in skills.items():
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, category.capitalize(), ln=True)
        pdf.set_font("Arial", "", 12)
        pdf.multi_cell(0, 10, ", ".join(sanitize_text(skill) for skill in skill_list))
    pdf.ln(5)
    
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Questions and Evaluations", ln=True)
    for i, q in enumerate(questions):
        if q['question'] in evaluations:
            data = evaluations[q['question']]
            evaluation = data["evaluation"]
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, f"Question {i+1}: {sanitize_text(q['question'])}", ln=True)
            pdf.set_font("Arial", "", 12)
            pdf.multi_cell(0, 10, f"Answer: {sanitize_text(data['answer'])}")
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, f"Score: {evaluation.get('score', 0)}/100", ln=True)
            pdf.set_font("Arial", "", 12)
            pdf.multi_cell(0, 10, f"Feedback: {sanitize_text(evaluation.get('feedback', 'No feedback'))}")
            missing = evaluation.get('missing_concepts', [])
            if missing:
                pdf.set_font("Arial", "B", 12)
                pdf.cell(0, 10, "Missing concepts:", ln=True)
                pdf.set_font("Arial", "", 12)
                for concept in missing:
                    pdf.cell(0, 10, f"- {sanitize_text(concept)}", ln=True)
            pdf.ln(5)
    
    output_path = f"interview_results_{candidate_name}_{interview_date.replace(':', '-')}.pdf"
    pdf.output(output_path)
    return output_path

# Firebase login
def login():
    if "user" not in st.session_state or "user_email" not in st.session_state:
        st.session_state.user = None
        st.session_state.user_email = None
    
    if not st.session_state.user:
        st.title("Login to Technical Interview Chatbot")
        login_option = st.selectbox("Choose login method", ["Email/Password", "Sign Up"])
        
        if login_option == "Email/Password":
            email = st.text_input("Email")
            password = st.text_input("Password", type="password")
            if st.button("Login"):
                try:
                    user = auth.sign_in_with_email_and_password(email, password)
                    st.session_state.user = user
                    st.session_state.user_email = user.get('email', None)
                    if not st.session_state.user_email:
                        st.error("No email found in user data. Please try again.")
                        return False
                    st.success("Logged in successfully!")
                    st.rerun()
                except Exception as e:
                    st.error(f"Login failed: {e}")
        
        elif login_option == "Sign Up":
            email = st.text_input("Email")
            password = st.text_input("Password", type="password")
            if st.button("Sign Up"):
                try:
                    user = auth.create_user_with_email_and_password(email, password)
                    st.session_state.user = user
                    st.session_state.user_email = user.get('email', None)
                    if not st.session_state.user_email:
                        st.error("No email found in user data. Please try again.")
                        return False
                    st.success("Account created and logged in!")
                    st.rerun()
                except Exception as e:
                    st.error(f"Sign up failed: {e}")
        
        return False
    return True

# Session state initialization
if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""
if "skills" not in st.session_state:
    st.session_state.skills = {}
if "questions" not in st.session_state:
    st.session_state.questions = []
if "current_question_index" not in st.session_state:
    st.session_state.current_question_index = 0
if "evaluations" not in st.session_state:
    st.session_state.evaluations = {}
if "interview_complete" not in st.session_state:
    st.session_state.interview_complete = False
if "interview_date" not in st.session_state:
    st.session_state.interview_date = datetime.now().strftime("%Y-%m-%d %H:%M")
if "chat_messages" not in st.session_state:
    st.session_state.chat_messages = [{"role": "assistant", "content": random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS)}]
if "candidate_name" not in st.session_state:
    st.session_state.candidate_name = ""
if "bot_state" not in st.session_state:
    st.session_state.bot_state = "wait_for_resume"
if "max_questions" not in st.session_state:
    st.session_state.max_questions = 5
if "debug_skills" not in st.session_state:
    st.session_state.debug_skills = []
if "raw_resume_text" not in st.session_state:
    st.session_state.raw_resume_text = ""
if "interview_history" not in st.session_state:
    st.session_state.interview_history = []
if "user" not in st.session_state:
    st.session_state.user = None
if "user_email" not in st.session_state:
    st.session_state.user_email = None

def add_message(role, content):
    st.session_state.chat_messages.append({"role": role, "content": content})

# Main app (behind login)
if login():
    # Sidebar
    with st.sidebar:
        st.header("Interview Settings")
        if st.session_state.bot_state in ["wait_for_resume", "analyzing_resume"]:
            st.info("Upload or paste your resume to begin.")
        elif st.session_state.bot_state == "interview":
            st.subheader("Progress")
            progress = st.session_state.current_question_index / len(st.session_state.questions)
            st.progress(progress)
            st.write(f"Question {st.session_state.current_question_index}/{len(st.session_state.questions)}")
            if st.session_state.skills:
                st.subheader("Skills Focus")
                for category, skills in st.session_state.skills.items():
                    with st.expander(category.capitalize()):
                        st.write(", ".join(skills))
        elif st.session_state.bot_state == "complete":
            st.success("Interview Complete!")
            evaluations = st.session_state.evaluations
            total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
            avg_score = total_score / len(evaluations) if evaluations else 0
            rating = "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement"
            st.metric("Overall Score", f"{avg_score:.1f}/100")
            st.metric("Rating", rating)
        
        max_q = st.slider("Number of Questions", min_value=3, max_value=10, value=st.session_state.max_questions)
        if max_q != st.session_state.max_questions:
            st.session_state.max_questions = max_q
        
        st.subheader("Interview History")
        if st.session_state.interview_history:
            for idx, record in enumerate(st.session_state.interview_history):
                with st.expander(f"Interview {idx+1}: {record['candidate_name']} ({record['date']})"):
                    st.write(f"**Average Score:** {record['avg_score']:.1f}/100")
                    st.write(f"**Rating:** {record['rating']}")
                    st.write("**Skills:**")
                    for category, skills in record['skills'].items():
                        st.write(f"- {category.capitalize()}: {', '.join(skills)}")
                    st.write("**Questions and Scores:**")
                    for i, q in enumerate(record['questions']):
                        if q['question'] in record['evaluations']:
                            eval_data = record['evaluations'][q['question']]
                            st.write(f"- Q{i+1}: {q['question']} (Score: {eval_data['evaluation'].get('score', 0)}/100)")
                    if st.button(f"Download PDF (Interview {idx+1})", key=f"download_{idx}"):
                        try:
                            pdf_path = export_results_as_pdf(record)
                            with open(pdf_path, "rb") as f:
                                pdf_bytes = f.read()
                            pdf_b64 = base64.b64encode(pdf_bytes).decode()
                            st.markdown(f'<a href="data:application/pdf;base64,{pdf_b64}" download="{os.path.basename(pdf_path)}">Download PDF</a>', unsafe_allow_html=True)
                        except Exception as e:
                            st.error(f"PDF generation error: {e}")
        
        if st.button("Start New Interview"):
            st.session_state.resume_text = ""
            st.session_state.skills = {}
            st.session_state.questions = []
            st.session_state.current_question_index = 0
            st.session_state.evaluations = {}
            st.session_state.interview_complete = False
            st.session_state.bot_state = "wait_for_resume"
            st.session_state.chat_messages = [{"role": "assistant", "content": random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS)}]
            st.session_state.candidate_name = ""
            st.session_state.debug_skills = []
            st.session_state.raw_resume_text = ""
            st.rerun()
        
        if st.button("Logout"):
            st.session_state.user = None
            st.session_state.user_email = None
            st.rerun()

    # Main app
    st.title("Technical Interview Chatbot 🤖")
    
    uploaded_file = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"], key="resume_upload")
    chat_container = st.container()
    
    def process_user_input(user_input):
        add_message("user", user_input)
        
        if st.session_state.bot_state == "wait_for_resume":
            if len(user_input.split()) <= 3 and len(st.session_state.chat_messages) <= 3:
                st.session_state.candidate_name = user_input
                add_message("assistant", f"Hi {user_input}! Please upload or paste your resume.")
                return
            
            st.session_state.resume_text = user_input
            st.session_state.bot_state = "analyzing_resume"
            add_message("assistant", "Analyzing your resume for skills...")
            
            skills = extract_skills(user_input)
            if not skills:
                add_message("assistant", "No technical skills found. List some skills (e.g., Python, Java, AWS).")
                st.session_state.bot_state = "manual_skills"
            else:
                st.session_state.skills = skills
                skill_message = random.choice(SKILL_MESSAGES) + "\n\n" + format_skills_message(skills)
                if st.session_state.debug_skills:
                    skill_message += "\n\n**Debug Info:**\n" + "\n".join(st.session_state.debug_skills)
                skill_message += "\n\nAre these correct? Add more skills or type 'start interview'."
                add_message("assistant", skill_message)
                st.session_state.bot_state = "confirm_skills"
        
        elif st.session_state.bot_state == "manual_skills":
            skills_input = user_input.lower()
            manual_skills = {}
            for category, skill_list in COMMON_SKILLS.items():
                found_skills = [skill for skill in skill_list if skill in skills_input]
                if found_skills:
                    manual_skills[category] = found_skills
            
            if not manual_skills:
                manual_skills = {'programming': ['python'], 'tools': ['git']}
            
            st.session_state.skills = manual_skills
            skill_message = "Added skills:\n\n" + format_skills_message(manual_skills) + "\n\nType 'start interview' to begin."
            add_message("assistant", skill_message)
            st.session_state.bot_state = "confirm_skills"
        
        elif st.session_state.bot_state == "confirm_skills":
            if any(x in user_input.lower() for x in ["start interview", "ready", "yes"]):
                technical_questions = generate_technical_questions(st.session_state.skills, st.session_state.max_questions)
                st.session_state.questions = technical_questions
                st.session_state.current_question_index = 0
                first_question = technical_questions[0]["question"] if technical_questions else "Tell me about your tech background."
                add_message("assistant", f"{random.choice(INTERVIEW_START_MESSAGES)}\n\n**Question 1:** {first_question}")
                st.session_state.bot_state = "interview"
            else:
                new_skills = extract_skills(user_input)
                if new_skills:
                    for category, skills_list in new_skills.items():
                        if category in st.session_state.skills:
                            st.session_state.skills[category].extend([s for s in skills_list if s not in st.session_state.skills[category]])
                        else:
                            st.session_state.skills[category] = skills_list
                    add_message("assistant", "Skills updated. Type 'start interview' to begin.")
                else:
                    add_message("assistant", "Type 'start interview' when ready.")
        
        elif st.session_state.bot_state == "interview":
            current_index = st.session_state.current_question_index
            current_question = st.session_state.questions[current_index]
            evaluation = evaluate_answer(
                question=current_question['question'],
                answer=user_input,
                expected_keywords=current_question['expected_keywords']
            )
            st.session_state.evaluations[current_question['question']] = {"answer": user_input, "evaluation": evaluation}
            
            score = evaluation.get('score', 0)
            feedback = f"{get_feedback_message(score)}\n\n**Score:** {score}/100\n\n{evaluation.get('feedback', '')}"
            missing = evaluation.get('missing_concepts', [])
            if missing:
                feedback += "\n\n**Improve:**\n" + "\n".join(f"- {concept}" for concept in missing)
            add_message("assistant", feedback)
            
            current_index += 1
            st.session_state.current_question_index = current_index
            
            if current_index < len(st.session_state.questions):
                next_question = st.session_state.questions[current_index]["question"]
                add_message("assistant", f"{random.choice(QUESTION_TRANSITIONS)}\n\n**Question {current_index + 1}:** {next_question}")
            else:
                st.session_state.bot_state = "complete"
                st.session_state.interview_complete = True
                evaluations = st.session_state.evaluations
                total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
                avg_score = total_score / len(evaluations) if evaluations else 0
                rating = "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement"
                
                interview_record = {
                    "candidate_name": st.session_state.candidate_name or "Candidate",
                    "date": st.session_state.interview_date,
                    "avg_score": avg_score,
                    "rating": rating,
                    "skills": st.session_state.skills,
                    "questions": st.session_state.questions,
                    "evaluations": st.session_state.evaluations
                }
                st.session_state.interview_history.append(interview_record)
                
                add_message("assistant", f"""
                ## Interview Complete!
                **Overall Score:** {avg_score:.1f}/100
                **Rating:** {rating}
                Options:
                1. Review answers
                2. Export PDF
                3. Send results via email
                4. View history
                5. Start new interview
                """)
        
        elif st.session_state.bot_state == "complete":
            if "review" in user_input.lower() or "answers" in user_input.lower():
                review = "## Your Responses\n\n"
                for i, q in enumerate(st.session_state.questions):
                    if q['question'] in st.session_state.evaluations:
                        data = st.session_state.evaluations[q['question']]
                        evaluation = data["evaluation"]
                        review += f"### Question {i+1}: {q['question']}\n"
                        review += f"**Answer:** {data['answer']}\n"
                        review += f"**Score:** {evaluation.get('score', 0)}/100\n"
                        review += f"**Feedback:** {evaluation.get('feedback', 'No feedback')}\n"
                        missing = evaluation.get('missing_concepts', [])
                        if missing:
                            review += "**Improve:**\n" + "\n".join(f"- {concept}" for concept in missing)
                        review += "\n---\n"
                add_message("assistant", review)
            
            elif "pdf" in user_input.lower() or "export" in user_input.lower():
                try:
                    evaluations = st.session_state.evaluations
                    total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
                    avg_score = total_score / len(evaluations) if evaluations else 0
                    interview_record = {
                        "candidate_name": st.session_state.candidate_name or "Candidate",
                        "date": st.session_state.interview_date,
                        "avg_score": avg_score,
                        "rating": "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement",
                        "skills": st.session_state.skills,
                        "questions": st.session_state.questions,
                        "evaluations": st.session_state.evaluations
                    }
                    pdf_path = export_results_as_pdf(interview_record)
                    with open(pdf_path, "rb") as f:
                        pdf_bytes = f.read()
                    pdf_b64 = base64.b64encode(pdf_bytes).decode()
                    add_message("assistant", f"PDF ready! <a href='data:application/pdf;base64,{pdf_b64}' download='{os.path.basename(pdf_path)}'>Download</a>")
                except Exception as e:
                    add_message("assistant", f"PDF generation error: {e}")
            
            elif "email" in user_input.lower() or "send" in user_input.lower():
                if st.session_state.user_email:
                    try:
                        evaluations = st.session_state.evaluations
                        total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
                        avg_score = total_score / len(evaluations) if evaluations else 0
                        interview_record = {
                            "candidate_name": st.session_state.candidate_name or "Candidate",
                            "date": st.session_state.interview_date,
                            "avg_score": avg_score,
                            "rating": "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement",
                            "skills": st.session_state.skills,
                            "questions": st.session_state.questions,
                            "evaluations": st.session_state.evaluations
                        }
                        pdf_path = export_results_as_pdf(interview_record)
                        if send_email(st.session_state.user_email, interview_record, pdf_path):
                            add_message("assistant", f"Results sent to {st.session_state.user_email}!")
                        else:
                            add_message("assistant", "Failed to send email. Try downloading the PDF.")
                    except Exception as e:
                        add_message("assistant", f"Error: {e}")
                else:
                    add_message("assistant", "No email available. Please log in again.")
            
            elif "history" in user_input.lower() or "past" in user_input.lower():
                if not st.session_state.interview_history:
                    add_message("assistant", "No past interviews. Complete one to build history!")
                else:
                    history_message = "## Interview History\n\n"
                    for idx, record in enumerate(st.session_state.interview_history):
                        history_message += f"### Interview {idx+1}: {record['candidate_name']} ({record['date']})\n"
                        history_message += f"**Score:** {record['avg_score']:.1f}/100\n"
                        history_message += f"**Rating:** {record['rating']}\n"
                        history_message += "**Skills:**\n"
                        for category, skills in record['skills'].items():
                            history_message += f"- {category.capitalize()}: {', '.join(skills)}\n"
                        history_message += "**Questions:**\n"
                        for i, q in enumerate(record['questions']):
                            if q['question'] in record['evaluations']:
                                eval_data = record['evaluations'][q['question']]
                                history_message += f"- Q{i+1}: {q['question']}\n"
                                history_message += f"  - **Answer:** {eval_data['answer']}\n"
                                history_message += f"  - **Score:** {eval_data['evaluation'].get('score', 0)}/100\n"
                        history_message += "\n---\n"
                    add_message("assistant", history_message)
            
            elif "new" in user_input.lower() or "start" in user_input.lower():
                st.session_state.resume_text = ""
                st.session_state.skills = {}
                st.session_state.questions = []
                st.session_state.current_question_index = 0
                st.session_state.evaluations = {}
                st.session_state.interview_complete = False
                st.session_state.bot_state = "wait_for_resume"
                st.session_state.chat_messages = [{"role": "assistant", "content": random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS)}]
                st.session_state.candidate_name = ""
                st.session_state.debug_skills = []
                st.session_state.raw_resume_text = ""
                st.rerun()
            else:
                add_message("assistant", """
                What's next?
                1. Review answers
                2. Export PDF
                3. Send results via email
                4. View history
                5. Start new interview
                """)
    
    # Handle file upload
    if uploaded_file is not None and not st.session_state.resume_text:
        file_extension = uploaded_file.name.split(".")[-1].lower()
        if file_extension == "pdf":
            resume_text = extract_text_from_pdf(uploaded_file)
        elif file_extension == "docx":
            resume_text = extract_text_from_docx(uploaded_file)
        else:
            resume_text = ""
            st.error("Please upload a PDF or DOCX file.")
        
        if resume_text:
            st.session_state.resume_text = resume_text
            st.session_state.bot_state = "analyzing_resume"
            st.session_state.chat_messages = [{"role": "assistant", "content": "Analyzing your resume..."}]
            skills = extract_skills(resume_text)
            if not skills:
                add_message("assistant", "No skills found. List some skills (e.g., Python, Java, AWS).")
                st.session_state.bot_state = "manual_skills"
            else:
                st.session_state.skills = skills
                skill_message = random.choice(SKILL_MESSAGES) + "\n\n" + format_skills_message(skills)
                if st.session_state.debug_skills:
                    skill_message += "\n\n**Debug:**\n" + "\n".join(st.session_state.debug_skills)
                skill_message += "\n\nCorrect? Add skills or type 'start interview'."
                add_message("assistant", skill_message)
                st.session_state.bot_state = "confirm_skills"
    
    # Chat interface
    with chat_container:
        for message in st.session_state.chat_messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"], unsafe_allow_html=True)
        
        if user_input := st.chat_input("Type here"):
            process_user_input(user_input)
            st.rerun()
    
    # Export on completion
    if st.session_state.interview_complete:
        with st.sidebar:
            st.subheader("Export")
            evaluations = st.session_state.evaluations
            total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
            avg_score = total_score / len(evaluations) if evaluations else 0
            if st.button("Generate PDF"):
                try:
                    interview_record = {
                        "candidate_name": st.session_state.candidate_name or "Candidate",
                        "date": st.session_state.interview_date,
                        "avg_score": avg_score,
                        "rating": "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement",
                        "skills": st.session_state.skills,
                        "questions": st.session_state.questions,
                        "evaluations": st.session_state.evaluations
                    }
                    pdf_path = export_results_as_pdf(interview_record)
                    with open(pdf_path, "rb") as f:
                        pdf_bytes = f.read()
                    pdf_b64 = base64.b64encode(pdf_bytes).decode()
                    st.markdown(f'<a href="data:application/pdf;base64,{pdf_b64}" download="{os.path.basename(pdf_path)}">Download PDF</a>', unsafe_allow_html=True)
                except Exception as e:
                    st.error(f"PDF generation error: {e}")
            
            st.subheader("Send Results")
            if st.button("Send Email"):
                if st.session_state.user_email:
                    try:
                        evaluations = st.session_state.evaluations
                        total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
                        avg_score = total_score / len(evaluations) if evaluations else 0
                        interview_record = {
                            "candidate_name": st.session_state.candidate_name or "Candidate",
                            "date": st.session_state.interview_date,
                            "avg_score": avg_score,
                            "rating": "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement",
                            "skills": st.session_state.skills,
                            "questions": st.session_state.questions,
                            "evaluations": st.session_state.evaluations
                        }
                        pdf_path = export_results_as_pdf(interview_record)
                        if send_email(st.session_state.user_email, interview_record, pdf_path):
                            st.success(f"Results sent to {st.session_state.user_email}!")
                        else:
                            st.error("Failed to send email. Try downloading the PDF.")
                    except Exception as e:
                        st.error(f"Email error: {e}")
                else:
                    st.error("No email available. Please log in again.")
